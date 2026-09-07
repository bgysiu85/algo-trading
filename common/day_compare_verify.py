#!/usr/bin/env python3
"""Check a built workbook against an independent re-computation.

    python -m common.day_compare_verify --xlsx var/reports/day_compare.xlsx \
        --summary var/reports/flex_symbol_days_all.csv \
        --round-trips var/reports/flex_round_trips.csv \
        --strategy mcl mc5 vw9_5m --reports var/reports/matched \
        --states var/state/matched

WHY THIS EXISTS
---------------
recalc.py answers "do the formulas evaluate". It cannot answer "are they
pointing at the right columns", and the difference is not academic: the first
build of this workbook recalculated 450 formulas with zero errors and every
summary total was wrong, because `trades` was written one column early and
shifted cost, gross and net one place left of their headers. A wrong SUMIFS
range is a clean, error-free file with wrong numbers, and nothing about reading
it says so.

So every summary cell is recomputed here from the same CSVs by a different code
path -- day_compare's own aggregation rather than the sheet's SUMIFS -- and
compared. Three classes of bug have been caught by exactly this:

  * the column shift above;
  * strategy trades on symbol-days absent from Ben's history, which sat in the
    clock sheet and not the monthly one, so two correct-looking sheets
    disagreed by $102.67;
  * day cells rounded to the cent before being SUMIFS'd, putting each month a
    few cents off the rows displayed directly above the total.

The cross-sheet check is the one that found the last two. Within one sheet
everything can be self-consistent and still wrong; the same money totalled
along three different axes has to land in the same place.

The workbook must have been recalculated first -- openpyxl writes formulas
with no cached values, so an unrecalculated file reads back as all None and
every check fails identically.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

from common import day_compare as D

FIELDS = ["days", "trades", "gross", "cost", "net"]
CLOCK_FIELDS = ["trades", "trades", "gross", "cost", "net"]
WIDTH = len(FIELDS)
FIRST_DATA_ROW = 6          # matches day_compare_xlsx's layout
TOL = 0.005                 # half a cent


def _read(wb, sheet, sources):
    ws = wb[sheet]
    out, r = {}, FIRST_DATA_ROW
    while True:
        k = ws.cell(row=r, column=1).value
        if k is None:
            break
        if k != "TOTAL":
            out[k] = [ws.cell(row=r, column=c).value
                      for c in range(2, 2 + len(sources) * WIDTH)]
        r += 1
    return out, ws, r - 1


def _total(wb, sheet, col):
    ws = wb[sheet]
    r = FIRST_DATA_ROW
    while ws.cell(row=r, column=1).value not in (None, "TOTAL"):
        r += 1
    return ws.cell(row=r, column=col).value


def compare(wb, sheet, buckets, sources, fields, fails):
    got, _ws, _last = _read(wb, sheet, sources)
    for key, cells in buckets.items():
        row = got.get(key)
        if row is None:
            fails.append(f"{sheet}: bucket {key!r} is missing from the sheet")
            continue
        for i, src in enumerate(sources):
            for j, f in enumerate(fields):
                want = getattr(cells[src], f)
                have = row[i * WIDTH + j]
                if have is None or abs(have - want) > TOL:
                    fails.append(f"{sheet} {key} {src} {f}: "
                                 f"sheet={have!r} recomputed={want!r}")
    return len(buckets)


def verify(xlsx: Path, summary: Path, round_trips: Path, reports: Path,
           states: Path, strategies: list[str]) -> list[str]:
    sources = ["mine"] + strategies
    wb = load_workbook(xlsx, data_only=True)
    rows, _scaled = D.build(summary, reports, states, strategies)
    units = {"mine": D.my_units(round_trips)}
    for n in strategies:
        units[n] = D.strategy_units(reports / f"backtest_trades_{n}.csv")
    units, _dropped = D.restrict_units(units, rows)

    fails: list[str] = []
    n_m = compare(wb, "By month", D.by_month(rows, sources), sources,
                  FIELDS, fails)
    n_w = compare(wb, "By weekday", D.by_weekday(rows, sources), sources,
                  FIELDS, fails)
    n_b = compare(wb, "By 30 min", D.by_entry_block(units, sources), sources,
                  CLOCK_FIELDS, fails)

    # THE CROSS-SHEET CHECK. Each sheet can be internally consistent and still
    # be summing a different population than its neighbour.
    for i, src in enumerate(sources):
        col = 2 + i * WIDTH + 4                       # the net column
        month = _total(wb, "By month", col)
        for other in ("By weekday", "By 30 min"):
            got = _total(wb, other, col)
            if month is None or got is None or abs(month - got) > TOL:
                fails.append(f"cross-sheet net for {src}: By month={month!r} "
                             f"vs {other}={got!r} -- the same money totalled "
                             "two ways must land in the same place")

    print(f"checked {n_m} months, {n_w} weekdays, {n_b} half-hour blocks "
          f"across {len(sources)} sources, plus cross-sheet totals")
    return fails


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Verify a day_compare workbook")
    ap.add_argument("--xlsx", default="var/reports/day_compare.xlsx")
    ap.add_argument("--summary", default="var/reports/flex_symbol_days_all.csv")
    ap.add_argument("--round-trips", default="var/reports/flex_round_trips.csv")
    ap.add_argument("--strategy", nargs="+", default=["mcl", "mc5", "vw9_5m"])
    ap.add_argument("--reports", default="var/reports")
    ap.add_argument("--states", default="var/state")
    a = ap.parse_args(argv)

    fails = verify(Path(a.xlsx), Path(a.summary), Path(a.round_trips),
                   Path(a.reports), Path(a.states), a.strategy)
    if not fails:
        print("OK -- every summary cell matches an independent recomputation")
        return 0
    print(f"\n{len(fails)} MISMATCH(ES):")
    for f in fails[:40]:
        print("  " + f)
    if len(fails) > 40:
        print(f"  ... and {len(fails) - 40} more")
    print("\nIf every cell is None, the workbook was not recalculated: "
          "openpyxl writes formulas with no cached values.")
    return 1


if __name__ == "__main__":
    sys.exit(main())
